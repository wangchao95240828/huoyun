#!/usr/bin/env python3
"""Generate docs/llm-api-pricing-comparison.xlsx from fixed rows. Requires openpyxl."""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:
    raise SystemExit("Install dependency: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "llm-api-pricing-comparison.xlsx"

ROWS = [
    (
        "厂商",
        "型号 / 档位",
        "输入\n(USD / 1M tokens)",
        "输出\n(USD / 1M tokens)",
        "备注",
    ),
    ("OpenAI", "GPT-5.5", 5.0, 30.0, "另有 cached input；Batch 约 -50%"),
    ("OpenAI", "GPT-5.4", 2.5, 15.0, "同上"),
    ("OpenAI", "GPT-5.4 mini", 0.75, 4.5, "同上"),
    ("Anthropic", "Claude Opus 4.7 / 4.6 / 4.5", 5.0, 25.0, "缓存分项；Batch 折价"),
    ("Anthropic", "Claude Sonnet 4.6 / 4.5 / 4", 3.0, 15.0, "同上"),
    ("Anthropic", "Claude Haiku 4.5", 1.0, 5.0, "同上"),
    ("Anthropic", "Claude Haiku 3", 0.25, 1.25, "同上"),
    ("Google Gemini", "3.1 Flash-Lite（Standard Paid）", 0.25, 1.5, "音频输入价更高；另有 Batch/Flex/Priority"),
    ("Google Gemini", "3.1 Pro Preview（≤200k 上下文）", 2.0, 12.0, "更长上下文档位更高"),
    ("MiniMax", "M2.7 / M2.5 / M2.1 等（标准）", 0.3, 1.2, "Prompt cache 读写另计价"),
    ("MiniMax", "*-highspeed*", 0.6, 2.4, "同上"),
    ("Kimi（月之暗面）", "各型号（国内开放平台）", None, None, "计价：¥ / 百万 tokens（CNY），见「官方链接」表"),
]

LINKS = [
    ("厂商", "官方定价 URL"),
    ("OpenAI", "https://openai.com/api/pricing/"),
    ("Anthropic Claude", "https://platform.claude.com/docs/en/about-claude/pricing"),
    ("Google Gemini", "https://ai.google.dev/gemini-api/docs/pricing"),
    ("MiniMax", "https://platform.minimax.io/docs/guides/pricing-paygo"),
    ("Kimi（Moonshot 国内）", "https://platform.moonshot.cn/"),
]

NOTES = [
    "【单位说明】1M tokens = 1 000 000 个 tokens；C/D 列数字为 USD/百万 tokens（除非该行备注写明 CNY）。",
    "输入 = 送进模型的 tokens；输出 = 模型生成 tokens（部分厂商含 reasoning，以官网为准）。",
    "说明：单价随模型、上下文、Batch/缓存/区域而变；以各厂商当天官网为准。",
    "整理时间参考：2026 年 5 月。",
    "",
    "对比注意：",
    "1）Tokenizer 不同，同样文本各家 token 计数不同；",
    "2）缓存、搜索/地图、音视频等可能单独计费；",
    "3）Kimi 请以控制台 ¥ / 百万 token 及具体模型 ID 为准。",
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "定价对照"

    unit_note = (
        "单位：以下 C、D 列为美元（USD）单价，计价粒度为每 1 000 000 个 tokens（1M = 百万）。"
        " Kimi 行为人民币计价，C/D 列勿填 USD。"
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    top = ws.cell(row=1, column=1, value=unit_note)
    top.alignment = Alignment(wrap_text=True, vertical="center")
    top.font = Font(bold=True)

    header_font = Font(bold=True)
    data_start = 2
    for r, row in enumerate(ROWS, start=data_start):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if r == data_start:
                cell.font = header_font
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c == 5 and val and r > data_start:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.row_dimensions[1].height = 36
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 40

    ws2 = wb.create_sheet("官方链接")
    for r, row in enumerate(LINKS, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws2.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = header_font
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 72

    ws3 = wb.create_sheet("说明")
    for r, line in enumerate(NOTES, start=1):
        ws3.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True)
    ws3.column_dimensions["A"].width = 88

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
